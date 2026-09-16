"""Annonce modifiable Word : contenu éditable, sans notes de fabrication.

Les codes-barres sont des images, présentes dans l'en-tête de première page
uniquement. Une modification du MRN nécessite leur régénération.
"""
from __future__ import annotations
import json
import os
import re
import tempfile
from pathlib import Path


def _texte(valeur):
    return str('à compléter' if valeur is None or valeur == '' else valeur).replace('à vérifier', 'à compléter')


def generer(data: dict, sortie: Path, dossier_data: Path | None = None):
    """Écriture atomique d'un DOCX éditable et de sa trace JSON séparée."""
    try:
        from docx import Document
        from docx.shared import Mm, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.graphics.barcode.code128 import Code128
        from . import pdfio
    except ImportError as exc:
        return False, f'Dépendance Word absente ({exc}) : relancer Installer.command ou Installer.cmd'

    def bordure(cellule, contour=False):
        props=cellule._tc.get_or_add_tcPr()
        for old in list(props.findall(qn('w:tcBorders'))): props.remove(old)
        borders=OxmlElement('w:tcBorders')
        for cote in ('top','left','bottom','right'):
            b=OxmlElement('w:'+cote);b.set(qn('w:val'),'single' if contour else 'nil')
            b.set(qn('w:sz'),'6');borders.append(b)
        props.append(borders)

    def paragraph(p, texte='', gras=False, taille=9):
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.line_spacing=1.05
        r=p.add_run(str(texte));r.bold=gras;r.font.size=Pt(taille)
        return p

    def table_largeurs(table, largeurs):
        table.autofit=False
        for col,w in zip(table.columns,largeurs):col.width=Mm(w)
        for row in table.rows:
            for cell,w in zip(row.cells,largeurs):
                cell.width=Mm(w);bordure(cell)
            trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))

    def titre_bloc(texte):
        p=doc.add_paragraph()
        paragraph(p,texte,True,10)
        p.paragraph_format.space_before=Pt(10)
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.keep_with_next=True
        props=p._p.get_or_add_pPr();b=OxmlElement('w:pBdr');line=OxmlElement('w:bottom')
        line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');line.set(qn('w:color'),'999999');b.append(line);props.append(b)

    def regler_table(t, largeurs):
        table_largeurs(t,largeurs)
        for i,row in enumerate(t.rows):
            for cell in row.cells:
                pr=cell._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for cote in ['top','bottom']:
                    e=OxmlElement('w:'+cote);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
                pr.append(margins)
                b=pr.find(qn('w:tcBorders'));bottom=b.find(qn('w:bottom'))
                bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'4');bottom.set(qn('w:color'),'BBBBBB')
                if i==0:
                    for para in cell.paragraphs:
                        for run in para.runs:run.bold=True

    def champs(lignes):
        if not lignes:return
        t=doc.add_table(rows=0,cols=2)
        for label,valeur in lignes:
            cells=t.add_row().cells
            paragraph(cells[0].paragraphs[0],str(label)+' :')
            paragraph(cells[1].paragraphs[0],_texte(valeur))
        table_largeurs(t,[46,114])
        doc.add_paragraph().paragraph_format.space_after=Pt(1)

    try:
        sortie.parent.mkdir(parents=True,exist_ok=True)
        doc=Document();section=doc.sections[0]
        section.page_width=Mm(210);section.page_height=Mm(297)
        section.top_margin=Mm(22);section.bottom_margin=Mm(22)
        section.left_margin=Mm(25);section.right_margin=Mm(25)
        section.header_distance=Mm(18);section.footer_distance=Mm(10)
        normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9)
        normal.font.color.rgb=RGBColor(0,0,0)
        normal.paragraph_format.space_after=Pt(3)
        for style in ('Title','Heading 1','Heading 2'):
            doc.styles[style].font.name='Arial';doc.styles[style].font.color.rgb=RGBColor(0,0,0)
        for style in doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        # Ce header n'est jamais répété sur les pages suivantes, même si Word repagine.
        section.different_first_page_header_footer=True
        header=section.first_page_header
        title=header.paragraphs[0];title.style=doc.styles['Title']
        paragraph(title,"Annonce d'arrivée",True,18)
        paragraph(header.add_paragraph(),'DM : '+_texte(data.get('dm'))).alignment=WD_ALIGN_PARAGRAPH.RIGHT
        sous='   '.join(f'{lib} : {_texte(data.get(cle))}' for cle,lib in [('type','Type'),('dossier','Dossier')])
        paragraph(header.add_paragraph(),sous,taille=8)
        if 'codes_barres' in data:
            mrns=[str(m) for m in data['codes_barres'] if re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9._/-]{0,79}',str(m))]
        else:
            mrns=[str(m) for m in data.get('mrns',[]) if re.fullmatch(r'\d{2}[A-Z]{2}[A-Z0-9]{14}',str(m))]
            if not mrns and re.fullmatch(r'\d{2}[A-Z]{2}[A-Z0-9]{14}',str(data.get('mrn',''))):
                mrns=[data['mrn']]
        if len(mrns)>10:raise ValueError('Maximum 10 MRN par annonce Word')
        with tempfile.TemporaryDirectory(prefix='.word-',dir=sortie.parent) as td:
            tmp=Path(td)
            if mrns:
                colonnes=min(3,len(mrns))
                grille=header.add_table(rows=(len(mrns)+colonnes-1)//colonnes,cols=colonnes,width=Mm(160))
                table_largeurs(grille,[160/colonnes]*colonnes)
                for i,mrn in enumerate(mrns):
                    largeur_mm=min(68,154/colonnes);largeur_pt=largeur_mm*72/25.4
                    pdf=tmp/f'code{i}.pdf';c=Canvas(str(pdf),pagesize=(largeur_pt,45))
                    barcode=Code128(mrn,barWidth=.48 if colonnes==3 else .72,barHeight=30,lquiet=10,rquiet=10)
                    barcode.drawOn(c,(largeur_pt-barcode.width)/2,13)
                    c.setFont('Helvetica',7);c.drawCentredString(largeur_pt/2,3,mrn);c.save()
                    png=pdfio.rendre_page_png(pdf,1,300,tmp,base=f'code{i}')
                    if not png:raise ValueError('Rendu du code-barres impossible')
                    cellule=grille.cell(i//colonnes,i%colonnes)
                    bordure(cellule);p=cellule.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(str(png),width=Mm(min(68,154/colonnes)))
            # Word réserve la hauteur de l'en-tête, sans répéter les codes-barres.
            titre_bloc('Références de transit')
            champs(data.get('entete') or [['MRN',data.get('mrn')]])
            if data.get('parties'):
                titre_bloc('Parties')
                groupes={}
                for pa in data['parties']:
                    role=re.sub(r'\s*\[[^]]+\]','',pa.get('role','Partie'))
                    lignes='\n'.join(pa.get('lignes',[]))
                    groupes.setdefault((role,lignes),[]).append(pa.get('role',''))
                for (role,lignes),origines in groupes.items():
                    refs=[m for origine in origines for m in re.findall(r'\[([^]]+)\]',origine)]
                    champs([(role,lignes)])
                    if refs:
                        paragraph(doc.add_paragraph(),'MRN : '+' / '.join(refs),taille=7)
            if data.get('transport'):
                titre_bloc('Transport et récapitulatif')
                par_mrn={};communs=[]
                for label,valeur in data['transport']:
                    match=re.fullmatch(r'(.*?)\s*\[([0-9A-Z]{18})\]',str(valeur))
                    if match and label in ('Nbre de colis','Masse brute totale'):
                        par_mrn.setdefault(match[2],{})[label]=match[1]
                    elif label != 'MRN du dépôt':
                        communs.append((label,valeur))
                if par_mrn:
                    recap=doc.add_table(rows=1,cols=3)
                    for c,texte in zip(recap.rows[0].cells,['MRN','Colis','Masse brute']):paragraph(c.paragraphs[0],texte,True,9)
                    for ref,valeurs in par_mrn.items():
                        for c,texte in zip(recap.add_row().cells,[ref,valeurs.get('Nbre de colis'),valeurs.get('Masse brute totale')]):paragraph(c.paragraphs[0],_texte(texte))
                    regler_table(recap,[90,30,40])
                    doc.add_paragraph().paragraph_format.space_after=Pt(2)
                champs(communs)
            p=doc.add_paragraph();p.paragraph_format.page_break_before=True
            p.style=doc.styles['Heading 1'];paragraph(p,"Liste d'inventaire",True,14)
            paragraph(doc.add_paragraph(),'Références : '+_texte(data.get('mrn')),taille=8).paragraph_format.space_after=Pt(14)
            t=doc.add_table(rows=1,cols=6)
            for cell,txt in zip(t.rows[0].cells,['Art.','Description','Code NC','Brut','Net','Colis']):
                paragraph(cell.paragraphs[0],txt,True,8)
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for a in data.get('articles',[]):
                row=t.add_row()
                designation=re.sub(r'\s*\[([^]]+)\]',r'\nMRN : \1',_texte(a.get('designation')))
                valeurs=[a.get('no'),designation,a.get('code_nc'),a.get('brut'),a.get('net'),a.get('colis')]
                for cell,val in zip(row.cells,valeurs):paragraph(cell.paragraphs[0],_texte(val),taille=8)
                for cle,label in [('marques','Marques'),('doc_prec','Doc. préc.'),('justif','Justif.')]:
                    if a.get(cle):paragraph(row.cells[1].add_paragraph(),f'{label} : {a[cle]}',taille=7)
            tot=data.get('total',{});cells=t.add_row().cells
            for cell,txt in zip(cells,['','TOTAL','',_texte(tot.get('brut')),_texte(tot.get('net')),_texte(tot.get('colis'))]):
                paragraph(cell.paragraphs[0],txt,True,8)
            regler_table(t,[9,62,25,22,22,20])
            for row in t.rows:
                for cell in row.cells[3:]:
                    for para in cell.paragraphs:para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph().paragraph_format.space_after=Pt(24)
            controle=doc.add_table(rows=1,cols=2);table_largeurs(controle,[72,88])
            gauche,droite=controle.rows[0].cells;ctl=data.get('controle',{})
            paragraph(gauche.paragraphs[0],'Dossier : '+_texte(ctl.get('dossier',data.get('dossier'))))
            paragraph(gauche.add_paragraph(),'DM : '+_texte(ctl.get('dm',data.get('dm'))))
            paragraph(gauche.add_paragraph(),'Réf. transit : '+_texte(ctl.get('ref_transit',data.get('mrn'))))
            bordure(droite,True)
            paragraph(droite.paragraphs[0],'CONTRÔLE : ULIX SWISS SA     '+str(ctl.get('numero','3140')),True,9)
            for label in ['CONFORME, selon document joint','NON-CONFORME',"AVIS D'IRRÉGULARITÉ ÉTABLI"]:
                paragraph(droite.add_paragraph(),'☐  '+label,taille=8).paragraph_format.space_after=Pt(8)
            paragraph(droite.add_paragraph(),'Date :                         Signature :',taille=8)
            for footer in (section.footer,section.first_page_footer):
                p=footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                paragraph(p,'ULIX SWISS SA     Page ',taille=7)
                f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
            temporaire=tmp/sortie.name;doc.save(temporaire)
            # Vérification structurelle : texte éditable et absence de protection.
            import zipfile
            from xml.etree import ElementTree as ET
            with zipfile.ZipFile(temporaire) as z:
                ET.fromstring(z.read('word/document.xml'))
                if b'documentProtection' in z.read('word/settings.xml'):
                    raise ValueError('Protection inattendue du document Word')
            trace=dossier_data or sortie.parent;trace.mkdir(parents=True,exist_ok=True)
            donnees=tmp/'trace.json';donnees.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
            os.replace(donnees,trace/(sortie.stem+'.json'))
            os.replace(temporaire,sortie)
        return True,str(sortie)
    except Exception as exc:
        return False,f'Génération Word impossible : {exc}'
